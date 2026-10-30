import streamlit as st
from openpyxl import load_workbook
from io import BytesIO
import pandas as pd


# --------------------------------------------------
# PAGE CONFIGURATION
# --------------------------------------------------

st.set_page_config(
    page_title="SLA Automation",
    page_icon="📊",
    layout="wide"
)

st.title("SLA Automation Dashboard")

st.write(
    "Upload your three Excel files to match incident numbers, "
    "populate SLA breach columns, and download the completed report."
)

st.divider()


# --------------------------------------------------
# CONSTANTS
# --------------------------------------------------

NUMBER_COL = "Number"
SOURCE_BREACHED_COL = "Has breached"

RESOLUTION_TARGET_COL = "Has Breached Resolution SLA"
RESPONSE_TARGET_COL = "Has Breached Response SLA"


# --------------------------------------------------
# HELPER FUNCTIONS
# --------------------------------------------------

def normalize_number(value):
    if value is None:
        return None

    value = str(value).strip()

    if not value or value.casefold() in {"nan", "none"}:
        return None

    return value


def normalize_header(value):
    if value is None:
        return ""

    return str(value).strip().casefold()


def find_column(headers, expected_name):
    expected = normalize_header(expected_name)

    for column_name, column_number in headers.items():
        if normalize_header(column_name) == expected:
            return column_number

    return None


def parse_boolean(value, filename, row_number):
    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)

    if value is not None:
        text = str(value).strip().casefold()

        if text == "true":
            return True

        if text == "false":
            return False

    raise ValueError(
        f"{filename}: Invalid or blank SLA value at Excel row "
        f"{row_number}. Expected TRUE or FALSE."
    )


def read_source_mapping(uploaded_file):
    workbook = load_workbook(
        BytesIO(uploaded_file.getvalue()),
        read_only=True,
        data_only=True
    )

    worksheet = workbook.worksheets[0]

    headers = {}

    for cell in worksheet[1]:
        if isinstance(cell.value, str):
            headers[cell.value.strip()] = cell.column

    number_column = find_column(headers, NUMBER_COL)
    breached_column = find_column(headers, SOURCE_BREACHED_COL)

    if number_column is None:
        workbook.close()
        raise ValueError(
            f"{uploaded_file.name}: Required column 'Number' was not found."
        )

    if breached_column is None:
        workbook.close()
        raise ValueError(
            f"{uploaded_file.name}: Required column "
            f"'Has breached' was not found."
        )

    mapping = {}

    for row_number in range(2, worksheet.max_row + 1):
        number = normalize_number(
            worksheet.cell(row_number, number_column).value
        )

        # Ignore rows without an incident number.
        if number is None:
            continue

        breached_value = parse_boolean(
            worksheet.cell(row_number, breached_column).value,
            uploaded_file.name,
            row_number
        )

        if number in mapping:
            if mapping[number] != breached_value:
                workbook.close()
                raise ValueError(
                    f"{uploaded_file.name}: Incident number '{number}' "
                    "has conflicting SLA breach values."
                )

            continue

        mapping[number] = breached_value

    workbook.close()

    return mapping


def process_files(imei_file, response_file, resolution_file):
    response_map = read_source_mapping(response_file)
    resolution_map = read_source_mapping(resolution_file)

    workbook = load_workbook(
        BytesIO(imei_file.getvalue())
    )

    worksheet = workbook.worksheets[0]

    headers = {}

    for cell in worksheet[1]:
        if isinstance(cell.value, str):
            headers[cell.value.strip()] = cell.column

    number_column = find_column(headers, NUMBER_COL)
    resolution_column = find_column(
        headers, RESOLUTION_TARGET_COL
    )
    response_column = find_column(
        headers, RESPONSE_TARGET_COL
    )

    if number_column is None:
        workbook.close()
        raise ValueError(
            "IMEI.xlsx: Required column 'Number' was not found "
            "in the first worksheet."
        )

    if resolution_column is None:
        workbook.close()
        raise ValueError(
            f"IMEI.xlsx: Required column '{RESOLUTION_TARGET_COL}' "
            "was not found in the first worksheet."
        )

    if response_column is None:
        workbook.close()
        raise ValueError(
            f"IMEI.xlsx: Required column '{RESPONSE_TARGET_COL}' "
            "was not found in the first worksheet."
        )

    total_rows = 0
    response_found = 0
    response_missing = 0
    resolution_found = 0
    resolution_missing = 0

    for row_number in range(2, worksheet.max_row + 1):
        number = normalize_number(
            worksheet.cell(row_number, number_column).value
        )

        if number is None:
            continue

        total_rows += 1

        # Populate the Resolution SLA column.
        if number in resolution_map:
            worksheet.cell(
                row_number, resolution_column
            ).value = resolution_map[number]

            resolution_found += 1

        else:
            worksheet.cell(
                row_number, resolution_column
            ).value = "NOT FOUND"

            resolution_missing += 1

        # Populate the Response SLA column.
        if number in response_map:
            worksheet.cell(
                row_number, response_column
            ).value = response_map[number]

            response_found += 1

        else:
            worksheet.cell(
                row_number, response_column
            ).value = "NOT FOUND"

            response_missing += 1

    output = BytesIO()
    workbook.save(output)
    workbook.close()

    output.seek(0)

    statistics = {
        "total_rows": total_rows,
        "response_found": response_found,
        "response_missing": response_missing,
        "resolution_found": resolution_found,
        "resolution_missing": resolution_missing,
    }

    return output.getvalue(), statistics


# --------------------------------------------------
# FILE UPLOAD SECTION
# --------------------------------------------------

st.subheader("Upload Excel Files")

col1, col2, col3 = st.columns(3)

with col1:
    imei_file = st.file_uploader(
        "1. Upload IMEI.xlsx",
        type=["xlsx"],
        key="imei"
    )

with col2:
    response_file = st.file_uploader(
        "2. Upload Response.xlsx",
        type=["xlsx"],
        key="response"
    )

with col3:
    resolution_file = st.file_uploader(
        "3. Upload Resolution.xlsx",
        type=["xlsx"],
        key="resolution"
    )


# --------------------------------------------------
# PROCESSING SECTION
# --------------------------------------------------

st.divider()

if st.button(
    "Process Excel Files",
    type="primary",
    use_container_width=True
):
    if not all([imei_file, response_file, resolution_file]):
        st.error("Please upload all three Excel files before processing.")

    else:
        try:
            with st.spinner("Processing Excel files..."):
                completed_file, stats = process_files(
                    imei_file,
                    response_file,
                    resolution_file
                )

            st.session_state["completed_file"] = completed_file
            st.session_state["statistics"] = stats
            st.session_state["output_filename"] = "Completed_IMEI.xlsx"

            st.success("Processing completed successfully!")

        except Exception as error:
            st.error(f"Processing failed: {error}")


# --------------------------------------------------
# RESULTS DASHBOARD
# --------------------------------------------------

if "statistics" in st.session_state:
    stats = st.session_state["statistics"]

    st.divider()
    st.subheader("Processing Summary")

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric(
            "Total Incident Rows",
            stats["total_rows"]
        )

    with col2:
        st.metric(
            "Response SLA Matches",
            stats["response_found"]
        )

    with col3:
        st.metric(
            "Resolution SLA Matches",
            stats["resolution_found"]
        )

    st.subheader("Matching Details")

    results_df = pd.DataFrame(
        {
            "Category": [
                "Response SLA",
                "Response SLA",
                "Resolution SLA",
                "Resolution SLA"
            ],
            "Status": [
                "Matched",
                "Not Found",
                "Matched",
                "Not Found"
            ],
            "Number of Rows": [
                stats["response_found"],
                stats["response_missing"],
                stats["resolution_found"],
                stats["resolution_missing"]
            ]
        }
    )

    st.dataframe(
        results_df,
        use_container_width=True,
        hide_index=True
    )

    if (
        stats["response_missing"] > 0
        or stats["resolution_missing"] > 0
    ):
        st.warning(
            "Some incident numbers were not found in the source files. "
            "Their corresponding SLA cells have been marked NOT FOUND."
        )

    else:
        st.success(
            "All incident numbers were matched in both source files."
        )

    st.divider()
    st.subheader("Download Completed Report")

    st.write(
        "Download the completed workbook. All existing IMEI columns "
        "are retained, and the two SLA columns are updated."
    )

    st.download_button(
        label="Download Completed_IMEI.xlsx",
        data=st.session_state["completed_file"],
        file_name=st.session_state["output_filename"],
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary",
        use_container_width=True
    )